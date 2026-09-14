"""从转换规则 Excel 直接导入 memory_rule，无需启动 Web 服务。"""
from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .db import Base, engine, SessionLocal
from .models import MemoryRule
from .normalize import normalize_value


@dataclass(frozen=True)
class RuleRecord:
    customer_name: str
    field_name: str
    old_value: str
    new_value: str
    source: str


@dataclass
class ImportPlan:
    records: list[RuleRecord] = field(default_factory=list)
    skipped_missing_replacement: int = 0
    skipped_empty_source: int = 0
    skipped_conflicts: int = 0


@dataclass
class ImportResult:
    inserted: int = 0
    skipped_existing: int = 0


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _rows(sheet) -> list[dict[str, str]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_text(value) for value in values[0]]
    return [
        {headers[index]: _text(value) for index, value in enumerate(row) if index < len(headers)}
        for row in values[1:]
    ]


def _rule_key(customer_name: str, field_name: str, old_value: str) -> tuple[str, str, str]:
    return (normalize_value(customer_name), field_name.strip(), normalize_value(old_value))


def build_import_plan(path: str | Path) -> ImportPlan:
    """将约定的两个工作表转换成安全、可导入的长期规则。"""
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"未找到 Excel 文件：{source}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        plan = ImportPlan()
        records: list[RuleRecord] = []

        if "字段替换表" in workbook.sheetnames:
            for row in _rows(workbook["字段替换表"]):
                customer = row.get("Customer Name", "")
                field_name = row.get("字段名", "")
                old_value = row.get("沉淀值", "")
                new_value = row.get("替换值", "")
                if not new_value:
                    plan.skipped_missing_replacement += 1
                    continue
                if not customer or not field_name or not old_value:
                    plan.skipped_empty_source += 1
                    continue
                records.append(RuleRecord(customer, field_name, old_value, new_value, "explicit_config"))

        if "Sheet1" in workbook.sheetnames:
            mappings: dict[str, list[tuple[str, str]]] = defaultdict(list)
            for row in _rows(workbook["Sheet1"]):
                original = row.get("智眸原值", "")
                target = row.get("ePortal目标值", "")
                if not original or not target:
                    plan.skipped_empty_source += 1
                    continue
                mappings[original].append((original, target))
            for original, values in mappings.items():
                targets = {target for _, target in values}
                if len(targets) != 1:
                    plan.skipped_conflicts += len(values)
                    continue
                records.append(RuleRecord(
                    original, "Customer Name", original, next(iter(targets)), "customer_name_mapping"
                ))

        seen: set[tuple[str, str, str]] = set()
        for record in records:
            key = _rule_key(record.customer_name, record.field_name, record.old_value)
            if key not in seen:
                seen.add(key)
                plan.records.append(record)
        return plan
    finally:
        workbook.close()


def import_plan(db: Session, plan: ImportPlan, *, dry_run: bool) -> ImportResult:
    """导入规则；相同匹配键已存在时保留数据库原值。"""
    existing = {
        _rule_key(rule.customer_name, rule.field_name, rule.old_value)
        for rule in db.query(MemoryRule).all()
    }
    result = ImportResult()
    for record in plan.records:
        key = _rule_key(record.customer_name, record.field_name, record.old_value)
        if key in existing:
            result.skipped_existing += 1
            continue
        result.inserted += 1
        if not dry_run:
            db.add(MemoryRule(
                customer_name=record.customer_name,
                field_name=record.field_name,
                old_value=record.old_value,
                new_value=record.new_value,
                rule_type="permanent",
                status="enabled",
                effective_count=0,
                source=record.source,
                created_by="excel_import",
                updated_by="excel_import",
            ))
        existing.add(key)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="导入 ePortal 转换规则 Excel 到 memory_rule")
    parser.add_argument("--file", required=True, help="Excel 文件的绝对或相对路径")
    parser.add_argument("--dry-run", action="store_true", help="只校验和统计，不写数据库")
    args = parser.parse_args(argv)

    plan = build_import_plan(args.file)
    if args.dry_run:
        result = ImportResult(inserted=len(plan.records))
    else:
        Base.metadata.create_all(engine, tables=[MemoryRule.__table__])
        db = SessionLocal()
        try:
            result = import_plan(db, plan, dry_run=False)
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    mode = "预览" if args.dry_run else "已导入"
    print(
        f"{mode}：新增 {result.inserted} 条，跳过已有 {result.skipped_existing} 条；"
        f"缺少替换值 {plan.skipped_missing_replacement} 条，空原值/必要字段 {plan.skipped_empty_source} 条，"
        f"冲突映射 {plan.skipped_conflicts} 条。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
