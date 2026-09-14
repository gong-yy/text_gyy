"""T2 人工修改界面相关接口：订单列表/详情、保存回写、失败重发。"""
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..db import get_db
from ..adapters.eportal import EPortalError, search_customers
from ..logging_config import audit
from ..models import Order, User
from ..schemas import SaveChangesRequest
from ..services.order import BizError, order_detail, replace_t2_attachment, resend, save_changes, sync_t2_attachment
from ..services.zhimou_callback import resend_final_result
from .deps import get_t2_operator

router = APIRouter(prefix="/api/orders", tags=["orders"])


def _get_order(db: Session, order_id: int) -> Order:
    order = db.get(Order, order_id)
    if not order:
        raise HTTPException(status_code=404, detail=f"订单不存在：{order_id}")
    return order


@router.get("/customers/search")
def customer_search(
    q: str = Query("", max_length=100), user: User = Depends(get_t2_operator), db: Session = Depends(get_db)
):
    try:
        return {"items": search_customers(db, q)}
    except EPortalError as exc:
        raise HTTPException(status_code=502, detail=exc.args[0]) from exc


@router.post("/products/import")
async def import_products(file: UploadFile = File(...), user: User = Depends(get_t2_operator)):
    """按 ePortal 源码导入首个 Excel 工作表的 A/B/C 产品列（跳过标题行）。"""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="请选择包含产品数据的 Excel 文件")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            items = []
            for product_id, description, vendor_part_no, *_ in sheet.iter_rows(min_row=2, values_only=True):
                row = {
                    "product_id": "" if product_id is None else str(product_id).strip(),
                    "description": "" if description is None else str(description).strip(),
                    "PN": "" if vendor_part_no is None else str(vendor_part_no).strip(),
                }
                if any(row.values()):
                    items.append(row)
        finally:
            workbook.close()
    except Exception as exc:
        raise HTTPException(status_code=422, detail="产品导入文件不是有效的 Excel 文件") from exc
    return {"items": items}


@router.get("/lookup")
def lookup_order_by_intellisight_id(
    intellisight_id: str = Query(..., min_length=1),
    user: User = Depends(get_t2_operator),
    db: Session = Depends(get_db),
):
    """供 ePortal 用智眸订单号打开 T2 时，解析为 T 系统本地订单编号。"""
    order = db.query(Order).filter(Order.zhimou_task_id == intellisight_id).one_or_none()
    if not order:
        raise HTTPException(status_code=404, detail=f"未找到 intellisight_id 对应的本地订单：{intellisight_id}")
    return {
        "order_id": order.id,
        "intellisight_id": order.zhimou_task_id,
        "form_id": order.form_id,
    }


@router.get("")
def list_orders(
    status: str | None = Query(None),
    customer: str | None = Query(None),
    user: User = Depends(get_t2_operator),
    db: Session = Depends(get_db),
):
    q = db.query(Order)
    if status:
        q = q.filter(Order.status == status)
    if customer:
        q = q.filter(Order.customer_name.contains(customer))
    rows = q.order_by(Order.id.desc()).limit(200).all()
    return [
        {
            "order_id": o.id,
            "form_id": o.form_id,
            "customer_name": o.customer_name,
            "status": o.status,
            "version": o.version,
            "holder": None,
            "updated_at": o.updated_at.isoformat(sep=" ") if o.updated_at else None,
        }
        for o in rows
    ]


@router.get("/{order_id}")
def get_order(order_id: int, user: User = Depends(get_t2_operator), db: Session = Depends(get_db)):
    return order_detail(db, _get_order(db, order_id))


@router.post("/{order_id}/attachments/{slot_key}")
async def replace_attachment(order_id: int, slot_key: str, file: UploadFile = File(...),
                             user: User = Depends(get_t2_operator), db: Session = Depends(get_db)):
    """T2 附件替换：立即同步附件，但不传 stage=3、也不提交 ePortal 表单。"""
    try:
        attachment = replace_t2_attachment(db, _get_order(db, order_id), slot_key, {
            "filename": file.filename,
            "content_type": file.content_type,
            "content": await file.read(),
        })
        sync_t2_attachment(db, _get_order(db, order_id))
        audit("t2_attachment_replaced", order_id=order_id, operator=user.username,
              slot=slot_key, filename=attachment["filename"])
        return {"attachment": attachment, "attachments": _get_order(db, order_id).payload.get("attachments", [])}
    except BizError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.message) from exc


@router.post("/{order_id}/lock")
def lock_order(order_id: int, user: User = Depends(get_t2_operator), db: Session = Depends(get_db)):
    """兼容旧客户端的入口通知；T2 不再占用或拒绝订单编辑。"""
    _get_order(db, order_id)
    audit("t2_edit_opened", order_id=order_id, operator=user.username)
    return {"ok": True, "holder": None, "locking": False}


@router.post("/{order_id}/unlock")
def unlock_order(order_id: int, user: User = Depends(get_t2_operator), db: Session = Depends(get_db)):
    _get_order(db, order_id)
    return {"ok": True, "locking": False}


@router.post("/{order_id}/save")
def save_order(order_id: int, body: SaveChangesRequest, user: User = Depends(get_t2_operator),
               db: Session = Depends(get_db)):
    """T2 保存：记忆确认对话框选项 + 负反馈选项随修改提交，保存后同步回写 ePortal。"""
    try:
        audit("t2_save_requested", order_id=order_id, operator=user.username,
              changed_fields=",".join(sorted(body.changes)))
        result = save_changes(
            db,
            _get_order(db, order_id),
            user,
            changes=dict(body.changes),
            items=body.items,
            memory_choices=dict(body.memory_choices),
            change_reasons=dict(body.change_reasons),
            feedback_choices=dict(body.feedback_choices),
        )
        audit("t2_save_completed", order_id=order_id, operator=user.username,
              status=result.get("status"), form_id=result.get("form_id"))
        return result
    except BizError as exc:
        audit("t2_save_rejected", order_id=order_id, operator=user.username, status=exc.status_code, reason=exc.message)
        raise HTTPException(status_code=exc.status_code, detail=exc.message)


@router.post("/{order_id}/resend")
def resend_order(order_id: int, user: User = Depends(get_t2_operator), db: Session = Depends(get_db)):
    """回写失败后手动重发（T2 数据本地留存不丢失）。"""
    try:
        audit("t2_resend_requested", order_id=order_id, operator=user.username)
        result = resend(db, _get_order(db, order_id))
        audit("t2_resend_completed", order_id=order_id, operator=user.username, status=result.get("status"))
        return result
    except BizError as exc:
        audit("t2_resend_rejected", order_id=order_id, operator=user.username, status=exc.status_code, reason=exc.message)
        raise HTTPException(status_code=exc.status_code, detail=exc.message)


@router.post("/{order_id}/zhimou-callback/resend")
def resend_zhimou_callback(
    order_id: int,
    user: User = Depends(get_t2_operator),
    db: Session = Depends(get_db),
):
    """重新发送已记录的智眸最终建单结果，不会再次调用 ePortal。"""
    order = _get_order(db, order_id)
    try:
        ok = resend_final_result(db, order)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    db.refresh(order)
    return {"ok": ok, "callback": (order.payload or {}).get("zhimou_callback")}
